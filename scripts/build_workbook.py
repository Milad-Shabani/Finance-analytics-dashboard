"""
TelNova Communications - Finance Analytics
Excel Workbook Builder
=================================================
Assembles the final, formatted, formula-driven .xlsx deliverable from the
cached synthetic datasets produced by generate_finance_data.py and
forecast_engine.py.

Design choices mirror the companion HR Analytics project:
  - Raw data lives in native Excel Tables (ListObjects) so formulas can use
    structured references and the HTML dashboard (SheetJS) can read clean
    tabular ranges.
  - The KPI_Dashboard sheet uses live SUMIFS / AVERAGEIFS / formulas against
    those tables -- nothing here is a hardcoded number.
  - Forecast/model outputs (revenue & EBITDA forecast, cash flow forecast,
    capex plan, budget risk scores) are written as values because they are
    the output of a statistical model computed in Python -- methodology is
    documented in the README sheet and on each model sheet.

Author: Milad Shabani
"""

import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from datetime import datetime, date
import os

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_ROOT = os.path.dirname(SCRIPT_DIR)
BASE = SCRIPT_DIR + os.sep
OUT_PATH = os.path.join(PROJECT_ROOT, "data", "Finance_Analytics_Workbook.xlsx")

NAVY = "1B2A4A"
TEAL = "0E7C7B"
WHITE = "FFFFFF"

HEADER_FONT = Font(name="Calibri", size=11, bold=True, color=WHITE)
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
TITLE_FONT = Font(name="Calibri", size=16, bold=True, color=NAVY)
SUBTITLE_FONT = Font(name="Calibri", size=10, italic=True, color="666666")
BODY_FONT = Font(name="Calibri", size=10)
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_header_row(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def autosize(ws, df, max_width=38):
    for i, col in enumerate(df.columns, start=1):
        series_len = df[col].apply(lambda x: len(str(x))).max() if len(df) else 0
        width = min(max(len(str(col)) + 2, (series_len or 0) + 2), max_width)
        ws.column_dimensions[get_column_letter(i)].width = width


def write_table(ws, df, start_row, table_name, style="TableStyleMedium9"):
    for j, col in enumerate(df.columns, start=1):
        ws.cell(row=start_row, column=j, value=col)
    for i, (_, row) in enumerate(df.iterrows(), start=start_row + 1):
        for j, col in enumerate(df.columns, start=1):
            val = row[col]
            if isinstance(val, (np.integer,)):
                val = int(val)
            elif isinstance(val, (np.floating,)):
                val = float(val)
            elif pd.isna(val):
                val = None
            elif isinstance(val, (pd.Timestamp, datetime)):
                val = val.date()
            ws.cell(row=i, column=j, value=val)

    n_rows = len(df)
    n_cols = len(df.columns)
    last_col_letter = get_column_letter(n_cols)
    last_row = start_row + n_rows
    ref = f"{get_column_letter(1)}{start_row}:{last_col_letter}{last_row}"
    table = Table(displayName=table_name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(name=style, showRowStripes=True, showFirstColumn=False)
    ws.add_table(table)
    autosize(ws, df)
    return start_row, last_row, n_cols


def add_cover_note(ws, title, subtitle, ncols=6):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(row=1, column=1, value=title)
    c.font = TITLE_FONT
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    c2 = ws.cell(row=2, column=1, value=subtitle)
    c2.font = SUBTITLE_FONT
    ws.row_dimensions[1].height = 26
    ws.row_dimensions[3].height = 6


def main():
    revenue = pd.read_csv(BASE + "_cache_revenue.csv")
    pnl = pd.read_csv(BASE + "_cache_pnl.csv")
    capex = pd.read_csv(BASE + "_cache_capex.csv")
    debt = pd.read_csv(BASE + "_cache_debt.csv")
    balance = pd.read_csv(BASE + "_cache_balance.csv")
    cashflow = pd.read_csv(BASE + "_cache_cashflow.csv")
    budget = pd.read_csv(BASE + "_cache_budget.csv")
    cust_econ = pd.read_csv(BASE + "_cache_customer_econ.csv")
    revenue_fc = pd.read_csv(BASE + "_cache_revenue_forecast.csv")
    cashflow_fc = pd.read_csv(BASE + "_cache_cashflow_forecast.csv")
    capex_fc = pd.read_csv(BASE + "_cache_capex_forecast.csv")
    budget_risk = pd.read_csv(BASE + "_cache_budget_risk.csv")

    wb = Workbook()
    wb.remove(wb.active)

    # ------------------------------------------------------------------
    # README
    # ------------------------------------------------------------------
    ws = wb.create_sheet("README")
    ws.sheet_view.showGridLines = False
    add_cover_note(ws, "TelNova Communications - Finance Analytics",
                    "Synthetic dataset generated for portfolio / demonstration purposes  |  Author: Milad Shabani", ncols=2)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 110

    rows = [
        ("Company profile", "Fictional converged telecom/ISP operator (same company as the companion HR Analytics project). ~1,047,500 subscribers, ~$255M annualized revenue, ~38% EBITDA margin."),
        ("Data status", "100% synthetically generated with a fixed random seed (reproducible). No real company's financials are represented."),
        ("3-statement consistency", "The Balance Sheet is built together with the Cash Flow Statement in a single pass so it balances EXACTLY (Assets = Liabilities + Equity, to the penny) -- cash is the model's plug, tied out through the cash flow statement, which is standard 3-statement-model practice."),
        ("Sheets - raw data", "Revenue_By_Segment, PnL, Balance_Sheet, Cash_Flow, Capex_Projects, Debt_Schedule, Budget_vs_Actual, Customer_Economics -- native Excel Tables, safe to filter/pivot directly."),
        ("Sheets - forecast", "Revenue_EBITDA_Forecast, Cashflow_Forecast, Capex_Plan: 12-month forward-looking financial plan. Methodology documented on each sheet."),
        ("Sheets - risk model", "Budget_Risk: an explainable weighted-factor model scoring every cost center 0-100 on continued budget-overrun risk. Methodology documented on the sheet."),
        ("Sheets - KPIs", "KPI_Dashboard: headline metrics computed with live SUM / AVERAGE / SUMIFS formulas referencing the raw data tables."),
        ("Currency", "All monetary figures in USD."),
        ("Companion dashboard", "dashboard/index.html reads this workbook's data (embedded at build time) to render an interactive finance dashboard -- see project README.md."),
    ]
    r = 4
    for label, text in rows:
        ws.cell(row=r, column=1, value=label).font = Font(bold=True, size=10, color=NAVY)
        cell = ws.cell(row=r, column=2, value=text)
        cell.font = BODY_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 32
        r += 1

    # ------------------------------------------------------------------
    # Revenue_By_Segment
    # ------------------------------------------------------------------
    ws = wb.create_sheet("Revenue_By_Segment")
    start, last_row, ncols = write_table(ws, revenue, 1, "tbl_Revenue")
    style_header_row(ws, 1, ncols)
    ws.freeze_panes = "A2"

    # ------------------------------------------------------------------
    # PnL
    # ------------------------------------------------------------------
    ws = wb.create_sheet("PnL")
    start, last_row, ncols = write_table(ws, pnl, 1, "tbl_PnL")
    style_header_row(ws, 1, ncols)
    ws.freeze_panes = "A2"

    chart = LineChart()
    chart.title = "Revenue vs EBITDA (48 months)"
    chart.style = 2
    chart.y_axis.title = "USD"
    rev_col = pnl.columns.get_loc("RevenueUSD") + 1
    ebitda_col = pnl.columns.get_loc("EBITDA_USD") + 1
    data = Reference(ws, min_col=rev_col, max_col=rev_col, min_row=1, max_row=last_row)
    data2 = Reference(ws, min_col=ebitda_col, max_col=ebitda_col, min_row=1, max_row=last_row)
    cats = Reference(ws, min_col=1, min_row=2, max_row=last_row)
    chart.add_data(data, titles_from_data=True)
    chart.add_data(data2, titles_from_data=True)
    chart.set_categories(cats)
    chart.width, chart.height = 26, 10
    ws.add_chart(chart, f"A{last_row + 3}")

    # ------------------------------------------------------------------
    # Balance_Sheet
    # ------------------------------------------------------------------
    ws = wb.create_sheet("Balance_Sheet")
    start, last_row, ncols = write_table(ws, balance, 1, "tbl_BalanceSheet")
    style_header_row(ws, 1, ncols)
    ws.freeze_panes = "A2"

    # ------------------------------------------------------------------
    # Cash_Flow
    # ------------------------------------------------------------------
    ws = wb.create_sheet("Cash_Flow")
    start, last_row, ncols = write_table(ws, cashflow, 1, "tbl_CashFlow")
    style_header_row(ws, 1, ncols)
    ws.freeze_panes = "A2"

    # ------------------------------------------------------------------
    # Capex_Projects
    # ------------------------------------------------------------------
    ws = wb.create_sheet("Capex_Projects")
    start, last_row, ncols = write_table(ws, capex, 1, "tbl_Capex")
    style_header_row(ws, 1, ncols)
    ws.freeze_panes = "A2"

    # ------------------------------------------------------------------
    # Debt_Schedule
    # ------------------------------------------------------------------
    ws = wb.create_sheet("Debt_Schedule")
    start, last_row, ncols = write_table(ws, debt, 1, "tbl_Debt")
    style_header_row(ws, 1, ncols)
    ws.freeze_panes = "A2"

    # ------------------------------------------------------------------
    # Budget_vs_Actual
    # ------------------------------------------------------------------
    ws = wb.create_sheet("Budget_vs_Actual")
    start, last_row, ncols = write_table(ws, budget, 1, "tbl_Budget")
    style_header_row(ws, 1, ncols)
    ws.freeze_panes = "A2"

    # ------------------------------------------------------------------
    # Customer_Economics
    # ------------------------------------------------------------------
    ws = wb.create_sheet("Customer_Economics")
    start, last_row, ncols = write_table(ws, cust_econ, 1, "tbl_CustEcon")
    style_header_row(ws, 1, ncols)

    # ------------------------------------------------------------------
    # Revenue_EBITDA_Forecast
    # ------------------------------------------------------------------
    ws = wb.create_sheet("Revenue_EBITDA_Forecast")
    add_cover_note(ws, "12-Month Revenue & EBITDA Forecast", "", ncols=6)
    ws.cell(row=3, column=1, value="Methodology: 50% linear-trend regression on last 18 months of revenue, blended 50% with a stated 8% annual growth target. EBITDA = forecast revenue x trailing-12-month average EBITDA margin.").font = SUBTITLE_FONT
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=6)
    start, last_row, ncols = write_table(ws, revenue_fc, 5, "tbl_RevenueForecast")
    style_header_row(ws, 5, ncols)

    fc_chart = LineChart()
    fc_chart.title = "Revenue & EBITDA Forecast (next 12 months)"
    fc_chart.y_axis.title = "USD"
    data = Reference(ws, min_col=2, max_col=4, min_row=5, max_row=last_row)
    data2 = Reference(ws, min_col=5, max_col=5, min_row=5, max_row=last_row)
    cats = Reference(ws, min_col=1, min_row=6, max_row=last_row)
    fc_chart.add_data(data, titles_from_data=True)
    fc_chart.add_data(data2, titles_from_data=True)
    fc_chart.set_categories(cats)
    fc_chart.width, fc_chart.height = 26, 10
    ws.add_chart(fc_chart, f"A{last_row + 3}")

    # ------------------------------------------------------------------
    # Cashflow_Forecast
    # ------------------------------------------------------------------
    ws = wb.create_sheet("Cashflow_Forecast")
    add_cover_note(ws, "12-Month Free Cash Flow & Liquidity Forecast", "", ncols=4)
    ws.cell(row=3, column=1, value="Methodology: Holt's linear (double) exponential smoothing (alpha=0.40, beta=0.25) fitted on 47 months of free cash flow, projected forward and cumulated against the latest cash balance.").font = SUBTITLE_FONT
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=6)
    start, last_row, ncols = write_table(ws, cashflow_fc, 5, "tbl_CashflowForecast")
    style_header_row(ws, 5, ncols)

    cf_chart = LineChart()
    cf_chart.title = "Projected Cash Balance (next 12 months)"
    data = Reference(ws, min_col=3, max_col=3, min_row=5, max_row=last_row)
    cats = Reference(ws, min_col=1, min_row=6, max_row=last_row)
    cf_chart.add_data(data, titles_from_data=True)
    cf_chart.set_categories(cats)
    cf_chart.width, cf_chart.height = 24, 9
    ws.add_chart(cf_chart, f"A{last_row + 3}")

    # ------------------------------------------------------------------
    # Capex_Plan
    # ------------------------------------------------------------------
    ws = wb.create_sheet("Capex_Plan")
    add_cover_note(ws, "12-Month Capital Investment Plan", "", ncols=4)
    ws.cell(row=3, column=1, value="Methodology: forecast revenue x trailing capex-intensity ratio, allocated to categories using each category's trailing-12-month share of actual spend. CapexType flags Growth (network expansion) vs. Maintenance.").font = SUBTITLE_FONT
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=6)
    start, last_row, ncols = write_table(ws, capex_fc, 5, "tbl_CapexPlan")
    style_header_row(ws, 5, ncols)

    # ------------------------------------------------------------------
    # Budget_Risk
    # ------------------------------------------------------------------
    ws = wb.create_sheet("Budget_Risk")
    add_cover_note(ws, "Cost Center Budget Overrun Risk Scoring", "", ncols=6)
    ws.cell(row=3, column=1, value="Methodology: explainable weighted-factor model (0-100). Weights: Avg. overspend 32%, Worsening trend 24%, Latest-month overspend 20%, Spend volatility 14%, Absolute budget size 10%. Min-max scaled across cost centers; bands set from this population's 50th/80th/93rd percentiles.").font = SUBTITLE_FONT
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=6)
    start, last_row, ncols = write_table(ws, budget_risk, 5, "tbl_BudgetRisk")
    style_header_row(ws, 5, ncols)

    score_col_idx = budget_risk.columns.get_loc("BudgetRiskScore") + 1
    score_col_letter = get_column_letter(score_col_idx)
    color_rule = ColorScaleRule(start_type="min", start_color="63BE7B", mid_type="percentile", mid_value=50,
                                 mid_color="FFEB84", end_type="max", end_color="F8696B")
    ws.conditional_formatting.add(f"{score_col_letter}6:{score_col_letter}{last_row}", color_rule)

    # ------------------------------------------------------------------
    # KPI_Dashboard (live formulas)
    # ------------------------------------------------------------------
    ws = wb.create_sheet("KPI_Dashboard", 0)
    ws.sheet_view.showGridLines = False
    add_cover_note(ws, "TelNova Communications - Finance KPI Dashboard (Live)", "All figures below are Excel formulas referencing tbl_PnL / tbl_BalanceSheet / tbl_CashFlow / tbl_Debt / tbl_Capex -- change the raw data and these recalculate automatically.", ncols=4)

    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 4
    ws.column_dimensions["D"].width = 42
    ws.column_dimensions["E"].width = 18

    n_pnl = len(pnl)
    kpis_left = [
        ("Latest Monthly Revenue (USD)", f"=INDEX(tbl_PnL[RevenueUSD],{n_pnl})", '"$"#,##0'),
        ("Annualized Revenue Run-Rate", f"=INDEX(tbl_PnL[RevenueUSD],{n_pnl})*12", '"$"#,##0'),
        ("Latest EBITDA Margin", f"=INDEX(tbl_PnL[EBITDA_MarginPct],{n_pnl})", "0.0%"),
        ("Latest Net Margin", f"=INDEX(tbl_PnL[NetMarginPct],{n_pnl})", "0.0%"),
        ("Trailing 12mo Avg EBITDA Margin", "=AVERAGE(OFFSET(tbl_PnL[EBITDA_MarginPct],COUNTA(tbl_PnL[Month])-12,0,12,1))", "0.0%"),
        ("Total Assets (latest)", f"=INDEX(tbl_BalanceSheet[TotalAssetsUSD],{n_pnl})", '"$"#,##0'),
        ("Cash & Equivalents (latest)", f"=INDEX(tbl_BalanceSheet[CashUSD],{n_pnl})", '"$"#,##0'),
        ("Total Debt (latest)", f"=INDEX(tbl_BalanceSheet[LongTermDebtUSD],{n_pnl})+INDEX(tbl_BalanceSheet[CurrentPortionDebtUSD],{n_pnl})", '"$"#,##0'),
        ("Net Debt / EBITDA (annualized)", f"=(INDEX(tbl_BalanceSheet[LongTermDebtUSD],{n_pnl})+INDEX(tbl_BalanceSheet[CurrentPortionDebtUSD],{n_pnl})-INDEX(tbl_BalanceSheet[CashUSD],{n_pnl}))/(INDEX(tbl_PnL[EBITDA_USD],{n_pnl})*12)", "0.00"),
        ("Current Ratio (latest)", f"=INDEX(tbl_BalanceSheet[CurrentAssetsUSD],{n_pnl})/INDEX(tbl_BalanceSheet[CurrentLiabilitiesUSD],{n_pnl})", "0.00"),
    ]

    kpis_right = [
        ("Total Capex, Trailing 12mo (USD)", "=SUM(OFFSET(tbl_Capex[ActualSpendUSD],COUNTA(tbl_Capex[ProjectID])-72,0,72,1))", '"$"#,##0'),
        ("Capex Intensity (Capex / Revenue, TTM)", "=SUM(OFFSET(tbl_Capex[ActualSpendUSD],COUNTA(tbl_Capex[ProjectID])-72,0,72,1))/(AVERAGE(OFFSET(tbl_PnL[RevenueUSD],COUNTA(tbl_PnL[Month])-12,0,12,1))*12)", "0.0%"),
        ("Free Cash Flow, Latest Month (USD)", "=INDEX(tbl_CashFlow[FreeCashFlowUSD],COUNTA(tbl_CashFlow[Month]))", '"$"#,##0'),
        ("Avg. Interest Rate Across Tranches", "=AVERAGE(tbl_Debt[InterestRatePct])", "0.00%"),
        ("Total Outstanding Debt (latest month)", f"=SUMIFS(tbl_Debt[OutstandingBalanceUSD],tbl_Debt[Month],INDEX(tbl_Debt[Month],COUNTA(tbl_Debt[Month])))", '"$"#,##0'),
        ("Avg. Monthly Budget Variance (24mo)", "=AVERAGE(tbl_Budget[VariancePct])", "0.0%"),
        ("Total Budgeted Spend (24mo, USD)", "=SUM(tbl_Budget[BudgetedUSD])", '"$"#,##0'),
        ("Total Actual Spend (24mo, USD)", "=SUM(tbl_Budget[ActualUSD])", '"$"#,##0'),
        ("Blended LTV:CAC Ratio (avg. across segments)", "=AVERAGE(tbl_CustEcon[LTV_CAC_Ratio])", "0.00"),
        ("Avg. Payback Period (months, across segments)", "=AVERAGE(tbl_CustEcon[PaybackPeriodMonths])", "0.0"),
    ]

    r0 = 5
    ws.cell(row=r0, column=1, value="PROFITABILITY & BALANCE SHEET").font = Font(bold=True, size=12, color=TEAL)
    ws.cell(row=r0, column=4, value="CAPITAL, DEBT & UNIT ECONOMICS").font = Font(bold=True, size=12, color=TEAL)
    for i, (label, formula, fmt) in enumerate(kpis_left):
        rr = r0 + 1 + i
        ws.cell(row=rr, column=1, value=label).font = BODY_FONT
        c = ws.cell(row=rr, column=2, value=formula)
        c.number_format = fmt
        c.font = Font(bold=True, size=11, color=NAVY)
    for i, (label, formula, fmt) in enumerate(kpis_right):
        rr = r0 + 1 + i
        ws.cell(row=rr, column=4, value=label).font = BODY_FONT
        c = ws.cell(row=rr, column=5, value=formula)
        c.number_format = fmt
        c.font = Font(bold=True, size=11, color=NAVY)

    note_row = r0 + max(len(kpis_left), len(kpis_right)) + 3
    ws.cell(row=note_row, column=1, value="Forecast, capital-plan and budget-risk figures live on their own sheets (Revenue_EBITDA_Forecast, Cashflow_Forecast, Capex_Plan, Budget_Risk) because they are statistical model output, not raw-data formulas.").font = SUBTITLE_FONT
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=5)

    order = ["KPI_Dashboard", "README", "Revenue_By_Segment", "PnL", "Balance_Sheet", "Cash_Flow",
             "Capex_Projects", "Debt_Schedule", "Budget_vs_Actual", "Customer_Economics",
             "Revenue_EBITDA_Forecast", "Cashflow_Forecast", "Capex_Plan", "Budget_Risk"]
    wb._sheets = [wb[name] for name in order]
    for name in order:
        wb[name].sheet_properties.tabColor = TEAL if name == "KPI_Dashboard" else NAVY

    wb.save(OUT_PATH)
    print(f"Workbook saved to {OUT_PATH}")


if __name__ == "__main__":
    main()
